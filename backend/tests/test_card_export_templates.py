from datetime import date, datetime
from io import BytesIO
from types import SimpleNamespace
from uuid import UUID, uuid4

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import create_engine, event, select
from sqlalchemy.dialects.postgresql import INET, JSONB
from sqlalchemy.ext.compiler import compiles, deregister
from sqlalchemy.orm import Session
from sqlalchemy.pool import StaticPool

from app.api.dependencies import get_actor_user_id, get_db_session
from app.domain.work_experience import WorkExperience, anchor_for_experience, format_work_experience
from app.main import create_app
from app.models import (
    AccessGrant,
    AuditEvent,
    Card,
    CardBlockInstance,
    CardEvent,
    CardEventChange,
    CardExportTemplate,
    CardTemplate,
    FieldValue,
    FormBlock,
    FormField,
    Organization,
    OrgUnit,
    Permission,
    ReferenceItem,
    ReferenceList,
    Registry,
    Role,
    User,
    role_permissions,
)
from app.models.base import Base


@pytest.fixture()
def export_context():
    compiles(JSONB, "sqlite")(lambda *_a, **_k: "JSON")
    compiles(INET, "sqlite")(lambda *_a, **_k: "TEXT")
    engine = create_engine(
        "sqlite://", connect_args={"check_same_thread": False}, poolclass=StaticPool
    )

    @event.listens_for(engine, "before_cursor_execute", retval=True)
    def adapt_json_defaults(_conn, _cursor, statement, parameters, _context, _many):
        return statement.replace("::jsonb", "").replace("::uuid", "").replace(
            "DEFAULT 'false'", "DEFAULT 0"
        ).replace("DEFAULT 'true'", "DEFAULT 1"), parameters

    with engine.connect() as connection:
        connection.connection.driver_connection.create_function(
            "gen_random_uuid", 0, lambda: uuid4().hex
        )
    Base.metadata.create_all(engine)
    with Session(engine, expire_on_commit=False) as session:
        admin = User(
            email="admin-test",
            display_name="Администратор",
            password_hash="unused",
            is_superuser=True,
        )
        viewer = User(email="viewer-test", display_name="Читатель", password_hash="unused")
        outsider = User(
            email="outsider-test", display_name="Другой читатель", password_hash="unused"
        )
        org = Organization(code="org", name="Организация А")
        child = Organization(code="child", name="Дочерняя организация")
        registry = Registry(code="test", name="Реестр")
        other_registry = Registry(code="other", name="Другой реестр")
        session.add_all([admin, viewer, outsider, org, child, registry, other_registry])
        session.flush()
        child.parent_id = org.id
        role = Role(code="reader", name="Читатель")
        permission = Permission(code="cards.read", description="Чтение карточек")
        session.add_all([role, permission])
        session.flush()
        session.execute(
            role_permissions.insert().values(role_id=role.id, permission_id=permission.id)
        )
        session.add(
            AccessGrant(
                user_id=viewer.id,
                role_id=role.id,
                organization_id=org.id,
                registry_id=registry.id,
                include_descendants=False,
            )
        )
        block = FormBlock(registry_id=registry.id, code="main", title="Основные сведения")
        session.add(block)
        session.flush()
        fields = [
            FormField(block_id=block.id, code=code, label=label, field_type=kind, position=index)
            for index, (code, label, kind) in enumerate(
                [
                    ("fio", "ФИО", "text"),
                    ("position", "Должность", "text"),
                    ("unit", "Подразделение", "text"),
                    ("appointed", "Дата назначения", "date"),
                    ("basis", "Основание назначения", "text"),
                ]
            )
        ]
        session.add_all(fields)
        session.flush()
        template = CardTemplate(
            registry_id=registry.id,
            code="main",
            name="Основной шаблон",
            field_schema_json={"field_ids": [str(f.id) for f in fields]},
        )
        session.add(template)
        session.flush()
        ctx = SimpleNamespace(
            session=session,
            admin=admin,
            viewer=viewer,
            outsider=outsider,
            org=org,
            child=child,
            registry=registry,
            other_registry=other_registry,
            block=block,
            fields=fields,
            template=template,
        )
        ctx.card = add_card(ctx, 10, "Альфа", date(2026, 9, 1))
        session.commit()
        app = create_app()
        app.dependency_overrides[get_db_session] = lambda: session
        app.dependency_overrides[get_actor_user_id] = lambda: ctx.admin.id
        with TestClient(app) as client:
            ctx.client = client
            ctx.app = app
            yield ctx
    engine.dispose()
    deregister(JSONB)
    deregister(INET)


def add_card(ctx, ordinal, fio, appointed, *, organization=None, template=None):
    card = Card(
        id=UUID(int=(15 << 124) + ordinal),
        registry_id=ctx.registry.id,
        card_template_id=(template or ctx.template).id,
        organization_id=(organization or ctx.org).id,
        lifecycle_status="active",
    )
    ctx.session.add(card)
    ctx.session.flush()
    instance = CardBlockInstance(card_id=card.id, block_id=ctx.block.id, ordinal=0)
    ctx.session.add(instance)
    ctx.session.flush()
    for field, value in zip(
        ctx.fields, [fio, "Специалист", "Отдел А", appointed, "Приказ № 1"], strict=True
    ):
        ctx.session.add(
            FieldValue(
                card_id=card.id,
                block_instance_id=instance.id,
                field_id=field.id,
                **{"value_date" if field.field_type == "date" else "value_text": value},
            )
        )
    ctx.session.flush()
    return card


def template_payload(ctx, kind="card_list"):
    mapping = {
        "position_field_id": str(ctx.fields[1].id),
        "structural_unit_field_id": str(ctx.fields[2].id),
        "appointment_date_field_id": str(ctx.fields[3].id),
        "appointment_basis_field_id": str(ctx.fields[4].id),
    }
    return {
        "code": "export",
        "name": "Выгрузка",
        "export_kind": kind,
        "card_template_id": str(ctx.template.id),
        "configuration_json": (
            mapping
            if kind == "personnel_changes"
            else {
                "field_ids": [str(ctx.fields[2].id), str(ctx.fields[0].id)],
            }
        ),
    }


def create_template(ctx, kind="card_list", *, payload=None):
    response = ctx.client.post(
        f"/api/v1/registries/{ctx.registry.id}/card-export-templates",
        json=payload or template_payload(ctx, kind),
    )
    assert response.status_code == 201, response.text
    return response.json()


def download(ctx, template, **overrides):
    return ctx.client.post(
        f"/api/v1/card-export-templates/{template['id']}/download",
        json={"organization_ids": [str(ctx.org.id)], **overrides},
    )


def workbook(response):
    assert response.status_code == 200, response.text[:200] if response.status_code != 200 else ""
    assert (
        response.headers["content-type"]
        == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    assert ".xlsx" in response.headers["content-disposition"]
    return load_workbook(BytesIO(response.content))


def test_export_template_does_not_persist_organizations(export_context):
    ctx = export_context
    template = create_template(ctx)

    assert "organization_ids" not in template["configuration_json"]


def test_download_rejects_empty_or_duplicate_request_organizations(export_context):
    ctx = export_context
    template = create_template(ctx)
    path = f"/api/v1/card-export-templates/{template['id']}/download"

    assert ctx.client.post(path, json={}).status_code == 422
    assert (
        ctx.client.post(
            path,
            json={"organization_ids": [str(ctx.org.id), str(ctx.org.id)]},
        ).status_code
        == 422
    )


def test_download_ignores_legacy_saved_organizations(export_context):
    ctx = export_context
    add_card(ctx, 11, "Бета", date(2026, 9, 1), organization=ctx.child)
    template = create_template(ctx)
    stored = ctx.session.get(CardExportTemplate, UUID(template["id"]))
    stored.configuration_json = {
        **stored.configuration_json,
        "organization_ids": [str(ctx.child.id)],
    }
    ctx.session.flush()

    sheet = workbook(download(ctx, template, organization_ids=[str(ctx.org.id)])).active

    assert [row[2] for row in list(sheet.values)[1:]] == ["Альфа"]


def test_card_list_export_combines_request_organizations_without_organization_column(
    export_context,
):
    ctx = export_context
    add_card(ctx, 11, "Бета", date(2026, 9, 1), organization=ctx.child)
    add_card(ctx, 12, "Аарон", date(2026, 9, 1), organization=ctx.child)
    template = create_template(ctx)

    sheet = workbook(
        download(ctx, template, organization_ids=[str(ctx.org.id), str(ctx.child.id)])
    ).active

    assert sheet.title == "Карточки"
    assert list(sheet.values)[0] == ("№ п/п", "Подразделение", "ФИО")
    assert [row[2] for row in list(sheet.values)[1:]] == ["Аарон", "Альфа", "Бета"]


def test_personnel_export_creates_one_sheet_per_request_organization(export_context):
    ctx = export_context
    template = create_template(ctx, "personnel_changes")

    book = workbook(
        download(
            ctx,
            template,
            organization_ids=[str(ctx.org.id), str(ctx.child.id)],
            period_from="2026-09-01",
            period_to="2026-09-30",
        )
    )

    assert len(book.worksheets) == 2
    assert all(sheet["A1"].value == "Сведения о кадровых изменениях" for sheet in book.worksheets)


def test_personnel_export_sorts_hires_and_events_by_fio(export_context):
    ctx = export_context
    aaron = add_card(ctx, 11, "Аарон", date(2026, 9, 2))
    add_event(ctx, ctx.card, 101, "change", date(2026, 9, 1), "Первое изменение")
    add_event(ctx, aaron, 102, "change", date(2026, 9, 2), "Второе изменение")
    ctx.session.flush()

    sheet = workbook(
        download(
            ctx,
            create_template(ctx, "personnel_changes"),
            period_from="2026-09-01",
            period_to="2026-09-30",
        )
    ).active
    sections = {
        cell.value: cell.row
        for row in sheet
        for cell in row
        if cell.value in {"Вновь приняты", "Иные изменения"}
    }

    assert [sheet.cell(sections["Вновь приняты"] + row, 1).value for row in (2, 3)] == [
        "Аарон",
        "Альфа",
    ]
    assert [sheet.cell(sections["Иные изменения"] + row, 1).value for row in (2, 3)] == [
        "Аарон",
        "Альфа",
    ]


def test_card_list_export_formats_work_experience_as_one_russian_column(export_context):
    ctx = export_context
    experience = WorkExperience(days=3, months=2, years=5)
    field = FormField(
        block_id=ctx.block.id,
        code="experience",
        label="Стаж",
        field_type="work_experience",
        position=10,
    )
    ctx.session.add(field)
    ctx.session.flush()
    ctx.template.field_schema_json = {
        "field_ids": [*ctx.template.field_schema_json["field_ids"], str(field.id)]
    }
    instance = ctx.session.scalar(
        select(CardBlockInstance).where(CardBlockInstance.card_id == ctx.card.id)
    )
    ctx.session.add(
        FieldValue(
            card_id=ctx.card.id,
            block_instance_id=instance.id,
            field_id=field.id,
            value_json={"anchor_date": anchor_for_experience(experience, date.today()).isoformat()},
        )
    )
    ctx.session.flush()
    payload = template_payload(ctx)
    payload["configuration_json"]["field_ids"] = [str(ctx.fields[0].id), str(field.id)]
    template = create_template(ctx, payload=payload)

    sheet = workbook(download(ctx, template)).active

    assert list(sheet.values)[0] == ("№ п/п", "ФИО", "Стаж")
    assert sheet.cell(2, 3).value == format_work_experience(experience)


def test_export_template_crud_persists_order_and_audits_soft_archive(export_context):
    ctx = export_context
    template = create_template(ctx)
    template_id = UUID(template["id"])
    stored = ctx.session.get(CardExportTemplate, template_id)
    assert stored.configuration_json["card_template_id"] == str(ctx.template.id)
    assert template["card_template_id"] == str(ctx.template.id)
    changed = ctx.client.patch(
        f"/api/v1/card-export-templates/{template_id}",
        json={
            "name": "Изменённая выгрузка",
            "configuration_json": {
                "field_ids": [str(ctx.fields[0].id), str(ctx.fields[2].id)],
                "organization_ids": [str(ctx.org.id)],
            },
        },
    )
    assert changed.status_code == 200, changed.text
    ctx.session.expire_all()
    assert ctx.session.get(CardExportTemplate, template_id).name == "Изменённая выгрузка"
    listed = ctx.client.get(f"/api/v1/registries/{ctx.registry.id}/card-export-templates").json()[
        "items"
    ]
    assert [item["id"] for item in listed] == [str(template_id)]
    assert ctx.client.get(f"/api/v1/card-export-templates/{template_id}").json()[
        "configuration_json"
    ]["field_ids"] == [str(ctx.fields[0].id), str(ctx.fields[2].id)]
    assert ctx.client.delete(f"/api/v1/card-export-templates/{template_id}").status_code == 200
    assert ctx.session.get(CardExportTemplate, template_id).archived_at is not None
    assert (
        ctx.client.get(f"/api/v1/registries/{ctx.registry.id}/card-export-templates").json()[
            "items"
        ]
        == []
    )
    assert download(ctx, template).status_code == 400
    audits = ctx.session.scalars(
        select(AuditEvent)
        .where(AuditEvent.object_id == template_id)
        .order_by(AuditEvent.created_at, AuditEvent.id)
    ).all()
    assert sorted(event.action for event in audits) == ["archive", "create", "update"]
    assert all(event.actor_user_id == ctx.admin.id for event in audits)


def test_card_list_export_template_keeps_configured_field_order(export_context):
    ctx = export_context
    add_card(ctx, 2, "Бета", date(2026, 8, 31))
    add_card(ctx, 3, "Скрытая карточка", date(2026, 9, 1), organization=ctx.child)
    value = ctx.session.scalar(
        select(FieldValue).where(
            FieldValue.card_id == ctx.card.id, FieldValue.field_id == ctx.fields[2].id
        )
    )
    value.value_text = "=1+1"
    ctx.session.flush()
    template = create_template(ctx)
    ctx.app.dependency_overrides[get_actor_user_id] = lambda: ctx.viewer.id
    book = workbook(download(ctx, template))
    assert book.sheetnames == ["Карточки"]
    sheet = book.active
    assert list(sheet.values) == [
        ("№ п/п", "Подразделение", "ФИО"),
        (1, "'=1+1", "Альфа"),
        (2, "Отдел А", "Бета"),
    ]
    assert sheet["B2"].data_type == "s"


def add_event(ctx, card, key, kind, day, basis, *, changes=True):
    row = CardEvent(
        id=UUID(int=(15 << 124) + key),
        card_id=card.id,
        event_type=kind,
        occurred_on=day,
        basis_text=basis,
    )
    ctx.session.add(row)
    ctx.session.flush()
    if kind == "change" and changes:
        ctx.session.add(
            CardEventChange(
                card_event_id=row.id,
                field_id=ctx.fields[1].id,
                old_value_json={
                    "field": {"label": "Старая подпись", "type": "text"},
                    "value": "До",
                },
                new_value_json={
                    "field": {"label": "Старая подпись", "type": "text"},
                    "value": "После",
                },
            )
        )
    ctx.session.flush()


def test_personnel_export_contains_three_merged_sections_and_event_rows(export_context):
    ctx = export_context
    beta = add_card(ctx, 2, "Бета", date(2026, 9, 30))
    outside = add_card(ctx, 3, "Вне периода", date(2026, 8, 31))
    other = add_card(ctx, 4, "Другая организация", date(2026, 9, 10), organization=ctx.child)
    beta.lifecycle_status = "dismissed"
    add_event(ctx, beta, 103, "dismissal", date(2026, 9, 30), "Увольнение № 2")
    add_event(ctx, ctx.card, 102, "change", date(2026, 9, 1), "Изменение № 1")
    add_event(ctx, outside, 101, "change", date(2026, 8, 31), "Нельзя включать")
    add_event(ctx, other, 104, "dismissal", date(2026, 9, 1), "Нельзя раскрывать")
    ctx.fields[1].label = "Новая подпись"
    ctx.session.flush()
    book = workbook(
        download(
            ctx,
            create_template(ctx, "personnel_changes"),
            period_from="2026-09-01",
            period_to="2026-09-30",
        )
    )
    assert len(book.sheetnames) == 1
    sheet = book.active
    assert sheet["A1"].value == "Сведения о кадровых изменениях"
    assert sheet["A2"].value == "Организация А"
    assert sheet["A3"].value == "За период с 01.09.2026 по 30.09.2026"
    assert sheet.max_column == 4
    # A short report must fit the printable landscape page, without orphaning a section.
    assert sum(sheet.row_dimensions[row].height for row in range(1, sheet.max_row + 1)) <= 450
    assert {"A1:D1", "A2:D2", "A3:D3"} <= {str(r) for r in sheet.merged_cells.ranges}
    sections = {
        cell.value: cell.row
        for row in sheet
        for cell in row
        if cell.value in {"Вновь приняты", "Уволены", "Иные изменения"}
    }
    assert list(sections) == ["Вновь приняты", "Уволены", "Иные изменения"]
    for row in sections.values():
        assert f"A{row}:D{row}" in {str(r) for r in sheet.merged_cells.ranges}
        assert sheet.cell(row + 1, 1).value == "Фамилия, имя, отчество"
        assert sheet.cell(row + 1, 1).border.bottom.style == "thin"
        assert sheet.cell(row + 2, 4).border.right.style == "thin"
    hired = sections["Вновь приняты"] + 2
    assert [sheet.cell(hired + i, 1).value for i in range(2)] == ["Альфа", "Бета"]
    assert sheet.cell(hired, 2).value == "Специалист\nОтдел А"
    assert sheet.cell(hired, 4).value == "01.09.2026, Приказ № 1"
    assert sheet.cell(hired - 1, 4).value == "Дата, основание назначения"
    assert {f"B{hired - 1}:C{hired - 1}", f"B{hired}:C{hired}"} <= {
        str(r) for r in sheet.merged_cells.ranges
    }
    dismissed = sections["Уволены"] + 2
    assert [sheet.cell(dismissed, c).value for c in range(1, 5)] == [
        "Бета",
        "Специалист",
        datetime(2026, 9, 30),
        "Увольнение № 2",
    ]
    assert sheet.cell(dismissed, 3).number_format == "DD.MM.YYYY"
    changed = sections["Иные изменения"] + 2
    assert sheet.cell(changed, 2).value == "Старая подпись: До → После"
    assert sheet.cell(changed, 4).value == "01.09.2026, Изменение № 1"
    assert sheet.cell(changed - 1, 4).value == "Дата, основание изменений"
    assert f"B{changed}:C{changed}" in {str(r) for r in sheet.merged_cells.ranges}
    text = " ".join(str(cell.value) for row in sheet for cell in row)
    assert "Нельзя" not in text and "Другая организация" not in text and "Вне периода" not in text


@pytest.mark.parametrize(
    "broken",
    [
        "duplicate",
        "foreign_template",
        "foreign_field",
        "missing_mapping",
        "wrong_date",
        "fio_mapping",
        "unknown_kind",
        "no_fio",
        "unexportable",
        "repeatable",
    ],
)
def test_export_template_rejects_invalid_configuration_without_audit(export_context, broken):
    ctx = export_context
    payload = template_payload(
        ctx,
        "personnel_changes"
        if broken in {"missing_mapping", "wrong_date", "fio_mapping"}
        else "card_list",
    )
    config = payload["configuration_json"]
    if broken == "duplicate":
        config["field_ids"] *= 2
    elif broken == "foreign_template":
        ctx.template.registry_id = ctx.other_registry.id
    elif broken == "foreign_field":
        config["field_ids"] = [str(uuid4())]
    elif broken == "missing_mapping":
        config.pop("position_field_id")
    elif broken == "wrong_date":
        ctx.fields[3].field_type = "text"
    elif broken == "fio_mapping":
        config["fio_field_id"] = str(ctx.fields[0].id)
    elif broken == "unknown_kind":
        payload["export_kind"] = "csv"
    elif broken == "no_fio":
        ctx.fields[0].is_active = False
    elif broken == "unexportable":
        ctx.fields[2].is_exportable = False
    elif broken == "repeatable":
        ctx.block.is_repeatable = True
    ctx.session.flush()
    response = ctx.client.post(
        f"/api/v1/registries/{ctx.registry.id}/card-export-templates", json=payload
    )
    assert response.status_code in {400, 422}, response.text
    assert ctx.session.scalars(select(CardExportTemplate)).all() == []
    assert ctx.session.scalars(select(AuditEvent)).all() == []


@pytest.mark.parametrize("period", [{}, {"period_from": "2026-10-01", "period_to": "2026-09-01"}])
def test_personnel_export_requires_valid_inclusive_period(export_context, period):
    ctx = export_context
    response = download(ctx, create_template(ctx, "personnel_changes"), **period)
    assert response.status_code == 400
    assert "период" in response.json()["detail"].lower()


def test_export_template_permissions_rechecked_on_every_operation(export_context):
    ctx = export_context
    template = create_template(ctx)
    ctx.app.dependency_overrides[get_actor_user_id] = lambda: ctx.viewer.id
    path = f"/api/v1/card-export-templates/{template['id']}"
    assert ctx.client.patch(path, json={"name": "Запрещено"}).status_code == 403
    assert ctx.client.delete(path).status_code == 403
    assert (
        ctx.client.post(
            f"/api/v1/registries/{ctx.registry.id}/card-export-templates",
            json=template_payload(ctx),
        ).status_code
        == 403
    )
    assert download(ctx, template, organization_ids=[str(ctx.child.id)]).status_code == 403
    assert download(ctx, template).status_code == 200
    ctx.fields[2].is_exportable = False
    ctx.session.flush()
    assert download(ctx, template).status_code == 400
    ctx.app.dependency_overrides[get_actor_user_id] = lambda: ctx.outsider.id
    assert download(ctx, template).status_code == 403
    assert ctx.client.get(path).status_code == 403
    assert (
        ctx.client.get(f"/api/v1/registries/{ctx.registry.id}/card-export-templates").status_code
        == 403
    )


def test_personnel_empty_sections_keep_bordered_placeholder_rows(export_context):
    ctx = export_context
    sheet = workbook(
        download(
            ctx,
            create_template(ctx, "personnel_changes"),
            period_from="2025-01-01",
            period_to="2025-01-31",
        )
    ).active
    empty_rows = [row[0].row for row in sheet if row[0].value == "Нет данных"]
    assert len(empty_rows) == 3
    for row in empty_rows:
        assert sheet.cell(row, 4).border.right.style == "thin"


def test_personnel_uses_unfilled_fio_label_and_keeps_long_change_readable(export_context):
    ctx = export_context
    fio_value = ctx.session.scalar(
        select(FieldValue).where(
            FieldValue.card_id == ctx.card.id, FieldValue.field_id == ctx.fields[0].id
        )
    )
    fio_value.value_text = None
    add_event(ctx, ctx.card, 107, "change", date(2026, 9, 1), "Основание\n" * 8)
    ctx.session.flush()
    sheet = workbook(
        download(
            ctx,
            create_template(ctx, "personnel_changes"),
            period_from="2026-09-01",
            period_to="2026-09-30",
        )
    ).active
    change_row = (
        next(cell.row for row in sheet for cell in row if cell.value == "Иные изменения") + 2
    )
    assert sheet.cell(change_row, 1).value == "Не заполнено"
    assert sheet.row_dimensions[change_row].height >= 120


def test_personnel_event_ties_are_stable_and_archived_cards_are_excluded(export_context):
    ctx = export_context
    from datetime import UTC

    beta = add_card(ctx, 2, "Бета", date(2026, 8, 1))
    archived = add_card(ctx, 3, "Архив", date(2026, 9, 1))
    archived.archived_at = datetime.now(UTC)
    add_event(ctx, ctx.card, 110, "change", date(2026, 9, 1), "Третий")
    add_event(ctx, beta, 109, "change", date(2026, 9, 1), "Второй")
    add_event(ctx, beta, 108, "change", date(2026, 9, 1), "Первый")
    add_event(ctx, archived, 111, "change", date(2026, 9, 1), "Архивное событие")
    template = create_template(ctx, "personnel_changes")
    rows = list(
        workbook(
            download(ctx, template, period_from="2026-09-01", period_to="2026-09-01")
        ).active.values
    )
    change_start = next(i for i, row in enumerate(rows) if row[0] == "Иные изменения") + 2
    assert [row[3] for row in rows[change_start:]] == [
        "01.09.2026, Третий",
        "01.09.2026, Первый",
        "01.09.2026, Второй",
    ]
    assert not any("Архив" in str(cell) for row in rows for cell in row)


@pytest.mark.parametrize("hidden", ["sensitive", "admin_block"])
def test_export_template_revalidates_field_visibility_for_reader(export_context, hidden):
    ctx = export_context
    template = create_template(ctx)
    if hidden == "sensitive":
        ctx.fields[2].sensitivity_level = "restricted"
    else:
        ctx.block.is_admin_only = True
    ctx.session.flush()
    ctx.app.dependency_overrides[get_actor_user_id] = lambda: ctx.viewer.id
    assert download(ctx, template).status_code == 403


def test_export_template_failed_patch_preserves_original_configuration(export_context):
    ctx = export_context
    template = create_template(ctx)
    response = ctx.client.patch(
        f"/api/v1/card-export-templates/{template['id']}",
        json={"name": "Не сохранять", "configuration_json": {"field_ids": [str(uuid4())]}},
    )
    assert response.status_code == 400
    ctx.session.expire_all()
    stored = ctx.session.get(CardExportTemplate, UUID(template["id"]))
    assert stored.name == "Выгрузка"
    assert stored.configuration_json["field_ids"] == template["configuration_json"]["field_ids"]
    assert [event.action for event in ctx.session.scalars(select(AuditEvent))] == ["create"]


@pytest.mark.parametrize("kind", ["select", "organization_ref", "org_unit_ref"])
def test_export_template_resolves_reference_labels_without_raw_ids(export_context, kind):
    ctx = export_context
    field = ctx.fields[2]
    field.field_type = kind
    value = ctx.session.scalar(
        select(FieldValue).where(FieldValue.card_id == ctx.card.id, FieldValue.field_id == field.id)
    )
    value.value_text = None
    if kind == "select":
        source = ReferenceList(code="units", name="Подразделения", registry_id=ctx.registry.id)
        ctx.session.add(source)
        ctx.session.flush()
        item = ReferenceItem(list_id=source.id, code="one", label="Доступное подразделение")
        ctx.session.add(item)
        ctx.session.flush()
        field.options_source_type = "reference_list"
        field.options_source_id = source.id
        value.value_reference_item_id = item.id
        expected = "Доступное подразделение"
    elif kind == "org_unit_ref":
        unit = OrgUnit(
            organization_id=ctx.org.id,
            code="unit",
            name="Доступное подразделение",
            type="department",
        )
        ctx.session.add(unit)
        ctx.session.flush()
        value.value_org_unit_id = unit.id
        expected = "Доступное подразделение"
    else:
        value.value_organization_id = ctx.child.id
        expected = "Недоступное значение"
    ctx.session.flush()
    template = create_template(ctx)
    ctx.app.dependency_overrides[get_actor_user_id] = lambda: ctx.viewer.id
    sheet = workbook(download(ctx, template)).active
    assert sheet["B2"].value == expected


@pytest.mark.parametrize(
    "kind",
    [
        "select",
        "multi_select",
        "organization_ref",
        "org_unit_ref",
        "card_ref",
        "registry_ref",
        "user_ref",
    ],
)
def test_personnel_snapshot_references_resolve_labels_and_hide_unavailable_ids(
    export_context, kind
):
    ctx = export_context
    field = FormField(block_id=ctx.block.id, code="historical", label="Ссылка", field_type=kind)
    ctx.session.add(field)
    ctx.session.flush()
    missing = uuid4()
    actor = ctx.viewer
    if kind in {"select", "multi_select"}:
        source = ReferenceList(
            code="history", name="Исторический справочник", registry_id=ctx.registry.id
        )
        ctx.session.add(source)
        ctx.session.flush()
        item = ReferenceItem(list_id=source.id, code="available", label="Доступный вариант")
        ctx.session.add(item)
        ctx.session.flush()
        field.options_source_type = "reference_list"
        field.options_source_id = source.id
        old_value = str(item.id)
        expected = "Доступный вариант"
    elif kind == "organization_ref":
        old_value, missing, expected = str(ctx.org.id), ctx.child.id, "Организация А"
    elif kind == "org_unit_ref":
        unit = OrgUnit(
            organization_id=ctx.org.id, code="visible", name="Доступный отдел", type="department"
        )
        hidden_unit = OrgUnit(
            organization_id=ctx.child.id, code="hidden", name="Секретный отдел", type="department"
        )
        ctx.session.add_all([unit, hidden_unit])
        ctx.session.flush()
        old_value, missing, expected = str(unit.id), hidden_unit.id, "Доступный отдел"
    elif kind == "card_ref":
        hidden_card = add_card(
            ctx, 27, "Секретная карточка", date(2026, 1, 1), organization=ctx.child
        )
        old_value, missing, expected = str(ctx.card.id), hidden_card.id, "Альфа"
    elif kind == "registry_ref":
        old_value, missing, expected = str(ctx.registry.id), ctx.other_registry.id, "Реестр"
    else:
        old_value, expected, actor = str(ctx.viewer.id), "Читатель", ctx.admin
    old_snapshot = [old_value] if kind == "multi_select" else old_value
    new_snapshot = [old_value, str(missing)] if kind == "multi_select" else str(missing)
    event_row = CardEvent(
        card_id=ctx.card.id, event_type="change", occurred_on=date(2026, 9, 1), basis_text="Приказ"
    )
    ctx.session.add(event_row)
    ctx.session.flush()
    ctx.session.add(
        CardEventChange(
            card_event_id=event_row.id,
            field_id=field.id,
            old_value_json={
                "field": {"label": "Историческая подпись", "type": kind},
                "value": old_snapshot,
            },
            new_value_json={
                "field": {"label": "Историческая подпись", "type": kind},
                "value": new_snapshot,
            },
        )
    )
    ctx.session.flush()
    template = create_template(ctx, "personnel_changes")
    ctx.app.dependency_overrides[get_actor_user_id] = lambda: actor.id
    sheet = workbook(
        download(ctx, template, period_from="2026-09-01", period_to="2026-09-01")
    ).active
    text = "\n".join(str(cell.value) for row in sheet for cell in row if cell.value is not None)
    assert f"Историческая подпись: {expected} → " in text
    assert "Недоступное значение" in text
    assert old_value not in text and str(missing) not in text and "Секрет" not in text


@pytest.mark.parametrize(
    "kind,value,expected",
    [
        ("date", "2026-08-20", "20.08.2026"),
        ("datetime", "2026-08-20T12:34:00", "20.08.2026 12:34"),
        ("bool", True, "Да"),
        ("number", "12.5", "12.5"),
    ],
)
def test_personnel_snapshot_formats_typed_values(export_context, kind, value, expected):
    ctx = export_context
    add_event(ctx, ctx.card, 124, "change", date(2026, 9, 1), "Приказ")
    change = ctx.session.scalar(select(CardEventChange))
    change.old_value_json = {"field": {"label": "Значение", "type": kind}, "value": None}
    change.new_value_json = {"field": {"label": "Значение", "type": kind}, "value": value}
    ctx.session.flush()
    sheet = workbook(
        download(
            ctx,
            create_template(ctx, "personnel_changes"),
            period_from="2026-09-01",
            period_to="2026-09-01",
        )
    ).active
    assert any(cell.value == f"Значение:  → {expected}" for row in sheet for cell in row)


def test_personnel_long_single_line_basis_and_change_get_wrapped_row_height(export_context):
    ctx = export_context
    add_event(ctx, ctx.card, 130, "change", date(2026, 9, 1), "Основание решения " * 70)
    change = ctx.session.scalar(select(CardEventChange))
    change.new_value_json = {
        "field": {"label": "Изменение", "type": "text"},
        "value": "Подробное описание " * 70,
    }
    ctx.session.flush()
    sheet = workbook(
        download(
            ctx,
            create_template(ctx, "personnel_changes"),
            period_from="2026-09-01",
            period_to="2026-09-01",
        )
    ).active
    row = next(
        cell.row
        for line in sheet
        for cell in line
        if isinstance(cell.value, str) and "Подробное описание" in cell.value
    )
    assert sheet.row_dimensions[row].height >= 350
    assert all(dimension.height <= 409.5 for dimension in sheet.row_dimensions.values())
    reconstructed = (
        "".join(str(sheet.cell(index, 2).value or "") for index in range(row, sheet.max_row + 1))
        .replace("\n", "")
        .replace(" ", "")
    )
    assert "Подробноеописание" * 70 in reconstructed
